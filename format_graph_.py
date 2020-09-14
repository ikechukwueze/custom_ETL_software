from openpyxl import load_workbook
from openpyxl.chart import ScatterChart, Reference, Series

import os





def format_graph(chart_data, multiple=None):
    wb = load_workbook(chart_data)
    ws = wb['Graph']

    # dimensions of the excel data ===> A1:D6
    table_dimension = ws.dimensions

    #split dimensions to [A1, D6]
    table_dimension = table_dimension.split(':')
    print(table_dimension)
    
    col_row_list = []
    
    #do this to split table_dimensions into columns and rows ===> [A, 1, D, 6]
    for col_row in table_dimension:
        for x in col_row:
            if x.isnumeric():
                ind = col_row.index(x)
                col_row_list.append(col_row[:ind])
                col_row_list.append(col_row[ind:])
                break
    print(col_row_list)
    


    #use the ord method to convert letters to numbers ==> [A, 1, D, 6] = [1, 1, 4, 6]
    y = [ord(i.lower()) - 96 if i.isalpha() else int(i) for i in col_row_list]



    min_column = y[0]+1
    min_row = y[1]
    max_column = y[2]
    max_row = y[3]





    chart = ScatterChart()
    chart.title = "Graph"
    chart.style = 2
    chart.x_axis.title = 'Distance'
    chart.y_axis.title = 'RSSI'



    x_data = Reference(ws, min_col=min_column, min_row=min_row+1, max_row=max_row)
    
    for i in range(min_column+1, max_column+1):
        values = Reference(ws, min_col=i, min_row=min_row, max_row=max_row)
        series = Series(values, x_data, title_from_data=True)
        chart.series.append(series)

    ws.add_chart(chart, "G10")

    wb.save(chart_data)

    #os.startfile(chart_data)