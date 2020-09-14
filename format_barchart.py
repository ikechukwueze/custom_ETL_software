from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Series, Reference
import os
#excelfile = r'C:\Users\TYDACOMM_DT9\Music\automate\dt_result.xlsx'


def format_barchart(chart_data):
    wb = load_workbook(chart_data)
    ws = wb['Bar Chart']

    # dimensions of the excel data ===> A1:D6
    table_dimension = ws.dimensions

    #split dimensions to [A1, D6]
    table_dimensions = table_dimension.split(':')
    
    col_row_list = []
    
    #do this to split table_dimensions into columns and rows ===> [A, 1, D, 6]
    for col_row in table_dimensions:
        
        for x in col_row:
            if x.isnumeric():
                ind = col_row.index(x)
                col_row_list.append(col_row[:ind])
                col_row_list.append(col_row[ind:])
                break



    #use the ord method to convert letters to numbers ==> [A, 1, D, 6] = [1, 1, 4, 6]
    y = [ord(i.lower()) - 96 if i.isalpha() else int(i) for i in col_row_list]

    min_column = y[0]
    min_row = y[1]
    max_column = y[2]
    max_row = y[3]

    chart1 = BarChart()
    chart1.type = "col"
    chart1.style = 4
    chart1.title = "Bar Chart"
    chart1.y_axis.title = '%age Count'
    chart1.x_axis.title = 'RSSI'



    #data = Reference(ws, min_col=min_column+1, min_row=min_row, max_col=max_column, max_row=max_row)
    cats = Reference(ws, min_col=min_column, min_row=min_row+1, max_row=max_row)
    



    for i in range(min_column+2, max_column+1, 2):

        data = Reference(ws, min_col=i, min_row=min_row, max_col=i, max_row=max_row)
        chart1.add_data(data, titles_from_data=True)
    

    print('Appending Bar chart')

    #chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    chart1.shape = 1
    ws.add_chart(chart1, "G10")

    wb.save(chart_data)
    print('Bar chart appended')
    #os.startfile(chart_data)
    

#r = r'C:\EZE\Mentum\9MOBILE_NIGERIA_PROJECT\resources\PRE_DRIVE\09042020_9mobile_gsm_wcdma_lte_pre\export\rd.xlsx'
#format_barchart(r)