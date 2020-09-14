from openpyxl import load_workbook
from openpyxl.chart import ScatterChart, Reference, Series







def format_graph(chart_data, multiple=None):
    wb = load_workbook(chart_data)
    ws = wb['Graph']

    # dimensions of the excel data ===> A1:D6
    table_dimension = ws.dimensions

    #split dimensions to [A1, D6]
    table_dimension = table_dimension.split(':')

    
    col_row_list = []
    
    #do this to split table_dimensions into columns and rows ===> [A, 1, D, 6]
    for col_row in table_dimension:
        for x in col_row:
            if x.isnumeric():
                ind = col_row.index(x)
                col_row_list.append(col_row[:ind])
                col_row_list.append(col_row[ind:])
                break

    


    #use the ord method to convert letters to numbers ==> [A, 1, D, 6] = [1, 1, 4, 6]
    y = [ord(i.lower()) - 96 if i.isalpha() else int(i) for i in col_row_list]


    datasets = int((y[2]-1)/4)


    min_column = [y[0]+1+(data * 4) for data in range(datasets)]


    min_row = y[1]

    max_column = [i+3 for i in min_column]


    max_row = y[3]

    #datasets = int(max_column-1/4)
    #c = 4



    chart = ScatterChart()
    chart.title = "Graph"
    chart.style = 2
    chart.x_axis.title = 'Distance'
    chart.y_axis.title = 'RSSI'



    
    for data in range(datasets):
        print('Appending chart')

        x_data = Reference(ws, min_col=min_column[data], min_row=min_row+1, max_row=max_row)
        
        for i in range(min_column[data]+1, max_column[data]+1):
            values = Reference(ws, min_col=i, min_row=min_row, max_row=max_row)
            series = Series(values, x_data, title_from_data=True)
            chart.series.append(series)

    print('Chart appended')
    
    ws.add_chart(chart, "G10")

    wb.save(chart_data)

    #os.startfile(chart_data)